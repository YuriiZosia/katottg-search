# Yurii Code system (YCsys) © 2026. Код, що працює, а не існує.
import os
import glob
import json
from datetime import datetime
import time
import sys
import pandas as pd
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

# Палітра ANSI
C_CYAN = '\033[96m'
C_GREEN = '\033[92m'
C_YELLOW = '\033[93m'
C_MAGENTA = '\033[95m'
C_GRAY = '\033[90m'
C_WHITE = '\033[97m'
C_RESET = '\033[0m' # Скидання кольору

# Цей магічний рядок вмикає підтримку ANSI-кольорів у Windows CMD
os.system("")

def print_banner():
    banner_full = [
        r"      ██╗   ██╗ ██████╗███████╗██╗   ██╗███████╗",
        r"      ╚██╗ ██╔╝██╔════╝██╔════╝╚██╗ ██╔╝██╔════╝",
        r"       ╚████╔╝ ██║     ███████╗ ╚████╔╝ ███████╗",
        r"        ╚██╔╝  ██║     ╚════██║  ╚██╔╝  ╚════██║",
        r"         ██║   ╚██████╗███████║   ██║   ███████║",
        r"         ╚═╝    ╚═════╝╚══════╝   ╚═╝   ╚══════╝",
        r"                                                ",
        r"      ██████╗  █████╗ ██████╗ ███████╗███████╗██████╗ ",
        r"      ██╔══██╗██╔══██╗██╔══██╗██╔════╝██╔════╝██╔══██╗",
        r"      ██████╔╝███████║██████╔╝███████╗█████╗  ██████╔╝",
        r"      ██╔═══╝ ██╔══██║██╔══██╗╚════██║██╔══╝  ██╔══██╗",
        r"      ██║     ██║  ██║██║  ██║███████║███████╗██║  ██║",
        r"      ╚═╝     ╚═╝  ╚═╝╚═╝  ╚═╝╚══════╝╚══════╝╚═╝  ╚═╝"
    ]

    current_year = datetime.now().year
    signature = f"Yurii Code system (YCsys) © {current_year}. Код, що працює, а не існує."
    
    # 1. Ефект "Сканування"
    print(f"{C_GRAY}=" * 65 + C_RESET)
    
    for i, line in enumerate(banner_full):
        color = C_CYAN if i < 7 else C_GREEN
        sys.stdout.write(f"\r{C_WHITE}>>> {line}{C_RESET}")
        sys.stdout.flush()
        time.sleep(0.05)
        sys.stdout.write(f"\r    {color}{line}{C_RESET}\n")
        sys.stdout.flush()

    # 2. Поява підзаголовку
    title = f"\n{C_YELLOW}[ П А Р С Е Р   К А Т О Т Т Г ]      {C_RESET}\n"
    for char in title:
        sys.stdout.write(char)
        sys.stdout.flush()
        time.sleep(0.02)

    # 3. Підпис
    print(f"{C_MAGENTA}{signature:^65}{C_RESET}")
    print(f"{C_GRAY}=" * 65 + C_RESET)
    
    # 4. Індикатори завантаження
    loading_text = f"\n{C_GREEN}[+]{C_RESET} Ініціалізація системи"
    sys.stdout.write(loading_text)
    for _ in range(3):
        time.sleep(0.3)
        sys.stdout.write(".")
        sys.stdout.flush()
    print(f" {C_GREEN}OK{C_RESET}")
    print(f"{C_GREEN}[+]{C_RESET} Пошук книг Excel...\n")


# Фіксовані константи для JS
REGIONS = [
    "Вінницька", "Волинська", "Дніпропетровська", "Донецька", "Житомирська", 
    "Закарпатська", "Запорізька", "Івано-Франківська", "Київська", "Кіровоградська", 
    "Луганська", "Львівська", "Миколаївська", "Одеська", "Полтавська", 
    "Рівненська", "Сумська", "Тернопільська", "Харківська", "Херсонська", 
    "Хмельницька", "Черкаська", "Чернівецька", "Чернігівська", "Київ", 
    "Севастополь", "Автономна Республіка Крим"
]
HROM_TYPES = ["міська", "селищна", "сільська"]
NP_TYPES = ["м. ", "с-ще ", "с. "]

def get_index(lst, item):
    try:
        return lst.index(item)
    except ValueError:
        return item

# Функція для динамічного створення словників
def get_dynamic_index(val, lst, val_map):
    if val not in val_map:
        val_map[val] = len(lst)
        lst.append(val)
    return val_map[val]

def main():
    print("=" * 60)
    print(" YCsys KATOTTG Parser (Максимальна оптимізація)")
    print("=" * 60)

    files = glob.glob("*.xlsx") + glob.glob("*.csv")
    
    if not files:
        print("❌ Файли .xlsx або .csv не знайдено в поточній папці!")
        return

    print("Знайдені файли:")
    for i, file_name in enumerate(files):
        print(f"[{i}] {file_name}")

    choice = input("\nВведіть номер файлу для опрацювання: ")
    try:
        selected_file = files[int(choice)]
    except (ValueError, IndexError):
        print("❌ Невірний вибір. Вихід.")
        return

    print(f"\n⏳ Читання файлу '{selected_file}'...")
    
    if selected_file.endswith('.csv'):
        df = pd.read_csv(selected_file, dtype=str)
    else:
        df = pd.read_excel(selected_file, dtype=str)

    level1_names = {}
    level2_names = {}
    level3_names = {}
    hromada_types = {}

    print("🔍 Аналіз ієрархії та пошук адміністративних центрів громад...")
    
    # Прохід 1: Визначення центрів та ієрархії
    for _, row in df.iterrows():
        l1 = str(row.get('Перший рівень', '')).strip()
        l2 = str(row.get('Другий рівень', '')).strip()
        l3 = str(row.get('Третій рівень', '')).strip()
        cat = str(row.get('Категорія об’єкта', '')).strip()
        name = str(row.get('Назва об’єкта', '')).strip()
        
        if pd.isna(name) or name == 'nan' or not name:
            continue
            
        if cat in ['O', 'K']:
            level1_names[l1] = name
        elif cat == 'P':
            level2_names[l2] = name
        elif cat == 'H':
            level3_names[l3] = name
        elif cat in ['M', 'X', 'C']:
            if l3 and l3 not in hromada_types:
                if cat == 'M': hromada_types[l3] = 'міська'
                elif cat == 'X': hromada_types[l3] = 'селищна'
                elif cat == 'C': hromada_types[l3] = 'сільська'

    print("✅ Типи громад визначено!")
    print("⚙️ Формування масивів даних (з динамічними індексами)...")

    excel_data = []
    js_data = []
    
    # Динамічні списки для районів та громад
    districts_list = []
    districts_map = {}
    hromadas_list = []
    hromadas_map = {}

    # Прохід 2: Формування фінальних списків
    for _, row in df.iterrows():
        cat = str(row.get('Категорія об’єкта', '')).strip()
        name = str(row.get('Назва об’єкта', '')).strip()
        
        if pd.isna(name) or name == 'nan' or not name:
            continue

        if cat in ['M', 'X', 'C']:
            l1 = str(row.get('Перший рівень', '')).strip()
            l2 = str(row.get('Другий рівень', '')).strip()
            l3 = str(row.get('Третій рівень', '')).strip()
            
            obl = level1_names.get(l1, '')
            ray = level2_names.get(l2, '')
            hrom = level3_names.get(l3, '')
            h_type = hromada_types.get(l3, '')
            
            np_type_excel = "місто" if cat == "M" else ("селище" if cat == "X" else "село")
            excel_data.append([obl, ray, hrom, h_type, np_type_excel, name])

            # Отримуємо/створюємо динамічні індекси для Району та Громади
            ray_idx = get_dynamic_index(ray, districts_list, districts_map)
            hrom_idx = get_dynamic_index(hrom, hromadas_list, hromadas_map)
            
            obl_idx = get_index(REGIONS, obl)
            h_type_idx = get_index(HROM_TYPES, h_type)
            np_type_js = "м. " if cat == "M" else ("с-ще " if cat == "X" else "с. ")
            np_type_idx = get_index(NP_TYPES, np_type_js)
            
            # ТЕПЕР ЗАМІСТЬ ТЕКСТУ ЗАПИСУЮТЬСЯ ІНДЕКСИ ДЛЯ РАЙОНУ І ГРОМАДИ!
            js_data.append([obl_idx, ray_idx, hrom_idx, h_type_idx, np_type_idx, name])

        elif cat == 'K':
            excel_data.append([name, '', name + 'ська', 'міська', 'місто', name])
            
            obl_idx = get_index(REGIONS, name)
            ray_idx = get_dynamic_index('', districts_list, districts_map)
            hrom_idx = get_dynamic_index(name + 'ська', hromadas_list, hromadas_map)
            h_type_idx = get_index(HROM_TYPES, 'міська')
            np_type_idx = get_index(NP_TYPES, 'м. ')
            
            js_data.append([obl_idx, ray_idx, hrom_idx, h_type_idx, np_type_idx, name])

    current_date = datetime.now().strftime("%d-%m-%Y")
    output_excel = f"опрацьований кодифікатор {current_date}.xlsx"
    output_js = "data_katottg.js"
    
    print(f"💾 Запис у файл {output_excel}...")
    df_out = pd.DataFrame(excel_data, columns=['Область', 'Район в області', 'Громада', 'Тип громади', 'Тип населеного пункту', 'Населений пункт'])
    df_out.to_excel(output_excel, index=False, sheet_name='КАТОТТГ')

    wb = openpyxl.load_workbook(output_excel)
    ws = wb.active
    max_row = ws.max_row
    tab = Table(displayName="KATOTTG_Table", ref=f"A1:F{max_row}")
    style = TableStyleInfo(name="TableStyleMedium14", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    ws.add_table(tab)
    
    widths = {'A': 30, 'B': 25, 'C': 30, 'D': 15, 'E': 20, 'F': 35}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    wb.save(output_excel)

    print(f"💾 Запис у файл {output_js}...")
    with open(output_js, 'w', encoding='utf-8') as f:
        f.write('const regions = ' + json.dumps(REGIONS, ensure_ascii=False) + ';\n')
        f.write('const hromTypes = ' + json.dumps(HROM_TYPES, ensure_ascii=False) + ';\n')
        f.write('const npTypes = ' + json.dumps(NP_TYPES, ensure_ascii=False) + ';\n')
        # Додаємо наші згенеровані масиви
        f.write('const districts = ' + json.dumps(districts_list, ensure_ascii=False) + ';\n')
        f.write('const hromadas = ' + json.dumps(hromadas_list, ensure_ascii=False) + ';\n\n')
        
        f.write('const katottgData = ' + json.dumps(js_data, ensure_ascii=False, separators=(',', ':')) + ';')

    print("\n🎉 ГОТОВО! YCsys успішно завершив роботу.")
    print(f"Оброблено населених пунктів: {len(excel_data)}")
    print(f"Унікальних районів: {len(districts_list)}")
    print(f"Унікальних громад: {len(hromadas_list)}")

if __name__ == "__main__":
    print_banner()
    main()

    input("\n\nНатисніть Enter, щоб вийти...")